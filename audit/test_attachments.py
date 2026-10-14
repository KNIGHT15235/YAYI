"""Exercise actual attachment extraction, including formerly successful placeholders."""
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend' / 'app'))
from router.attachment_router import process_attachment


class AttachmentTests(unittest.TestCase):
    def extract(self, extension, create):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ('audit' + extension)
            create(path)
            attachment = SimpleNamespace(filename=path.name, status='pending')
            db = MagicMock()
            db.query.return_value.filter.return_value.first.return_value = attachment
            process_attachment('audit', str(path), lambda: db)
            return attachment

    def test_text(self):
        attachment = self.extract('.txt', lambda p: p.write_text('银杏42', encoding='utf-8'))
        self.assertEqual(attachment.content_text, '银杏42')
        self.assertEqual(attachment.status, 'completed')

    def test_docx_includes_tables(self):
        from docx import Document
        def create(path):
            document = Document()
            document.add_paragraph('正文银杏42')
            document.add_table(rows=1, cols=1).cell(0, 0).text = '表格内容'
            document.save(path)
        attachment = self.extract('.docx', create)
        self.assertEqual(attachment.status, 'completed')
        self.assertIn('正文银杏42', attachment.content_text)
        self.assertIn('表格内容', attachment.content_text)

    def test_xlsx(self):
        from openpyxl import Workbook
        def create(path):
            workbook = Workbook()
            workbook.active.append(['测试', 42])
            workbook.save(path)
            workbook.close()
        attachment = self.extract('.xlsx', create)
        self.assertEqual(attachment.status, 'completed')
        self.assertIn('测试\t42', attachment.content_text)

    def test_pdf(self):
        from PyPDF2 import PdfWriter
        from PyPDF2.generic import DecodedStreamObject, DictionaryObject, NameObject
        def create(path):
            writer = PdfWriter()
            page = writer.add_blank_page(width=300, height=200)
            page = writer.pages[0]
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            stream.set_data(b'BT /F1 12 Tf 20 100 Td (Audit text 42.) Tj ET')
            page[NameObject('/Contents')] = writer._add_object(stream)
            with path.open('wb') as target:
                writer.write(target)
        attachment = self.extract('.pdf', create)
        self.assertEqual(attachment.status, 'completed')
        self.assertIn('Audit text 42.', attachment.content_text)

    def test_unsupported_not_false_success(self):
        attachment = self.extract('.bin', lambda p: p.write_bytes(b'placeholder'))
        self.assertEqual(attachment.status, 'failed')
        self.assertIn('尚不支持', attachment.error_message)


if __name__ == '__main__':
    unittest.main()
